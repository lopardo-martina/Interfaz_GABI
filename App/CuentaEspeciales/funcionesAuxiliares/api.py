"""
App/CuentaEspeciales/funcionesAuxiliares/enviar_api.py
 
Lee el Excel unificado y envía los registros a la API en lotes de 10.000,
usando paralelismo para mayor velocidad.
 
Orden de ejecución: 4° dentro del post-proceso.
"""

import os
import requests
import numpy as np
import pandas as pd
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from openpyxl import load_workbook


from App.CuentaEspeciales.config import (
    ROW_API_DOMINO, ROW_API_TOKEN_URL,
    ROW_API_USUARIO, ROW_API_CONTRASENIA,
)


def obtener_token(token_url:str, username:str, password:str) -> str:
    response = requests.post(token_url, data={'username': username, 'password': password})
    print(response.status_code)
    print(response.text)
    if response.status_code == 200:
        return response.json().get('access')
    else:
        raise Exception(f"No se pudo obtener el token.")
    

def validar_fecha(fecha) -> str | None:
    """Devuelve la fecha en formato YYYY-MM-DD o None si no es válida."""
    
    formatos_validos = ['%Y-%m-%d', '%Y/%m/%d', '%d-%m-%Y', '%d/%m/%Y']
    
    try:
        if isinstance(fecha, str):
            # Intentar convertir el string a datetime usando formatos válidos
            for formato in formatos_validos:
                try:
                    return datetime.strptime(fecha, formato).strftime('%Y-%m-%d')
                except ValueError:
                    continue
                
        elif isinstance(fecha, datetime):
            # Si ya es datetime, formatearlo directamente
            return fecha.strftime('%Y-%m-%d')
        
    except ValueError:
        return None  # Fecha no válida
    return None


def convertir_fechas_a_string(registro:dict) -> dict:
    """Convierte los campos de fecha del registro a string ISO."""
    for k, v in registro.items():
        if v is None:
            continue
        if k in ['vencimiento', 'periodo_calculado']:
            if isinstance(v, str):
                f = pd.to_datetime(v, dayfirst=True, errors='coerce')
                registro[k] = f.strftime('%Y-%m-%d') if not pd.isnull(f) else None
            elif isinstance(v, datetime):
                registro[k] = v.strftime('%Y-%m-%d')
        elif isinstance(v, datetime):
            registro[k] = v.strftime('%Y-%m-%d')
    return registro


def limpiar_datos(df: pd.DataFrame) -> pd.DataFrame:
    """Reemplaza NaN/inf por None y convierte columnas datetime a string."""
    df = df.replace({np.nan: None, np.inf: None, -np.inf: None})
    for col in df.select_dtypes(include=[np.datetime64]).columns:
        df[col] = df[col].dt.strftime('%Y-%m-%d')
    if 'alta' in df.columns:
        df['alta'] = df['alta'].replace('NaT', None)
    return df


def excel_a_json(ruta_excel: str) -> tuple[pd.DataFrame, list]:
    """
    Leer el archivo Excel y convertirlo a JSON.
    """
    df = pd.read_excel(ruta_excel, dtype={'comercio': str, 'cuenta': str})
    df = limpiar_datos(df)
    
    # Agregar columna para el estado del envío
    df['estado_envio'] = None
    
    json_data = df.to_dict(orient='records')
    return df, json_data



def enviar_lote(lote: list, api_url: str, headers: dict, log, lote_num: int):
    """Envía un lote de registros en paralelo con ThreadPoolExecutor."""

    cont_exito = 0
    errores_totales = []

    def enviar_uno(reg):
        reg  = convertir_fechas_a_string(reg)
        resp = requests.post(api_url, json=reg, headers=headers)
        if resp.status_code == 201:
            return True, None
        return False, f"Error {resp.status_code}: {resp.text}"
 
    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(enviar_uno, r): r for r in lote}
        for fut in as_completed(futures):
            ok, err = fut.result()
            if ok:
                cont_exito += 1
            elif err:
                errores_totales.append(err)
 
    if cont_exito:
        log(f"Lote {lote_num}: {cont_exito} registros enviados", "ok")
    for err in errores_totales:
        log(f"Lote {lote_num}: {err}", "error")
 
    return cont_exito, errores_totales



def enviar_datos_a_api(outputs_dir: str, settings_path: str, log=None):
    """Lee el Excel unificado del día y lo envía a la API en lotes.
 
    Args:
        outputs_dir:   ruta a data/outputs/
        settings_path: ruta al settings.xlsx
        log:           función de logging — recibe (mensaje, tipo)
    """
    if log is None:
        log = lambda msg, tipo="": print(f"[{tipo.upper() or 'INFO'}] {msg}", flush=True)
 
    # Leer credenciales y URLs del settings.xlsx
    try:
        wb = load_workbook(settings_path)
        hoja = wb["Sheet1"]
        dominio = hoja.cell(row=ROW_API_DOMINO, column=2).value
        token_url = hoja.cell(row=ROW_API_TOKEN_URL, column=2).value
        usuario = hoja.cell(row=ROW_API_USUARIO, column=2).value
        contrasena = hoja.cell(row=ROW_API_CONTRASENIA, column=2).value
    except Exception as e:
        log(f"Error leyendo settings.xlsx: {e}", "error")
        return
 
    api_url = f"{dominio}/partidas/api/deudas/"
    
    #Busca el unificado.xlsx del dia
    fecha  = datetime.now()
    carpeta_dia = os.path.join(
        outputs_dir,
        f"{fecha.strftime('%Y')}",
        fecha.strftime("%m"),
        fecha.strftime("%d"),
    )
    ruta_unificado = os.path.join(carpeta_dia, "unificado.xlsx")
 
    if not os.path.exists(ruta_unificado):
        log(f"No se encontró el archivo unificado: {ruta_unificado}", "error")
        return
 
    log(f"Leyendo: {ruta_unificado}", "info")
    
    try:
        _, json_data = excel_a_json(ruta_unificado)
    except Exception as e:
        log(f"Error leyendo Excel: {e}", "error")
        return
 
    log(f"Total registros a enviar: {len(json_data)}", "info")
    
    
    #Enviar lotes de 10.000 registros
    cont_exito_total = 0
    errores_totales = []
    lote_num = 1
    
    # Dividir en lotes de 10000 registros
    for idx in range(0, len(json_data), 10000):
        lote = json_data[idx:idx + 10000]
        log(f"Enviando lote {lote_num} ({len(lote)} registros)...", "info")
        
        try:
            headers = {
                'Authorization': f'Bearer {obtener_token(token_url, usuario, contrasena)}',
                'Content-Type': 'application/json',
            }
            exito, errores = enviar_lote(lote, api_url, headers, log, lote_num)
            cont_exito_total += exito
            errores_totales.extend(errores)
        except Exception as e:
            log(f"Error en lote {lote_num}: {e}", "error")
            
        lote_num += 1
    
    log(f"Envío completado: {cont_exito_total} registros exitosos, {len(errores_totales)} errores", "ok")




if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..")))
    from App.CuentaEspeciales.config import OUTPUTS_DIR, SETTINGS_PATH
    enviar_datos_a_api(OUTPUTS_DIR, SETTINGS_PATH)