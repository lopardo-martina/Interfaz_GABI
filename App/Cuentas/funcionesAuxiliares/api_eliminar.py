"""
App/CuentaEspeciales/funcionesAuxiliares/eliminar_deudas.py
 
Carga las cuentas del Excel de entrada y las elimina en la API
antes de subir los datos nuevos.
 
Orden de ejecución: 3° dentro del post-proceso.
"""
 
 
import json
import os
import requests
import openpyxl
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook

from App.CuentaEspeciales.config import (
    ROW_API_DOMINO, ROW_API_TOKEN_URL,
    ROW_API_USUARIO, ROW_API_CONTRASENIA,
    INPUTS_DIR, SETTINGS_PATH,
    ROW_NOMBRE_EXCEL
)

   

def obtener_token(token_url, username, password):
    response = requests.post(token_url, data={'username': username, 'password': password})
    print(response.status_code)
    print(response.text)
    if response.status_code == 200:
        return response.json().get('access')
    else:
        raise Exception(f"No se pudo obtener el token.")
    
    

def cargar_cuentas_desde_excel(ruta_archivo:str, log) -> list:
    """
    Lee las cuentas del Excel y las devuelve como lista
    formateadas a 20 dígitos con ceros a la izquierda.
    """
    try:
        df = pd.read_excel(ruta_archivo, dtype={'cuenta': str})
        # Suponiendo que la columna con las cuentas se llama 'cuenta'
        cuentas = df['cuenta'].dropna().astype(str)
        # Formatear las cuentas a 20 dígitos con ceros a la izquierda
        cuentas = cuentas.apply(lambda x: x.zfill(20)).tolist()
        return cuentas
    except Exception as e:
        log(f"Error al cargar el archivo Excel: {e}", "error")
        return []


# Función para eliminar deudas
def eliminar_deudas(inputs_dir: str, settings_path: str, log=None):
    """
    Elimina las deudas de las cuentas en la API.
 
    """
    if log is None:
        log = lambda msg, tipo="": print(f"[{tipo.upper() or 'INFO'}] {msg}", flush=True)
 
    # Leer credenciales y URLs del settings.xlsx
    try:
        wb = load_workbook(settings_path)
        hoja = wb["Sheet1"]
        dominio = hoja.cell(row=ROW_API_DOMINO, column=2).value
        usuario = hoja.cell(row=ROW_API_USUARIO, column=2).value
        contrasena = hoja.cell(row=ROW_API_CONTRASENIA, column=2).value
        token_url = hoja.cell(row=ROW_API_TOKEN_URL, column=2).value
        nombre_excel = hoja.cell(row=ROW_NOMBRE_EXCEL, column=2).value
    except Exception as e:
        log(f"Error leyendo settings.xlsx: {e}", "error")
        return
 
    API_URL = f"{dominio}/partidas/api/deudas/"
    
    #Cargar cuentas
    ruta_archivo = os.path.join(inputs_dir, f"{nombre_excel}.xlsx")
    cuentas = cargar_cuentas_desde_excel(ruta_archivo, log)
    if not cuentas:
        log("No se encontraron cuentas para eliminar", "warn")
        return
    
    # LLAmar la API
    try:
        headers = {
                'Authorization': f'Bearer {obtener_token(token_url, usuario, contrasena)}',
                'Content-Type': 'application/json',
            }
        data = {"cuentas": cuentas}

        print('Eliminar deuda de: \n', cuentas)

        response = requests.post(API_URL, headers=headers, json=data)
        
        if response.status_code == 200:
            print(response.json().get("message"))
        else:
            log(f"Error al eliminar deudas: {response.status_code} - {response.text}", "error")
        
    except Exception as e:
        log(f"Error en la API: {e}", "error")


# Ejecución del script
if __name__ == "__main__":
    import sys
    sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..")))
    from App.CuentaEspeciales.config import INPUTS_DIR, SETTINGS_PATH
    eliminar_deudas(INPUTS_DIR, SETTINGS_PATH)
